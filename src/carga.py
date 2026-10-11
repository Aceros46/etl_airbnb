"""
carga.py
--------
Fase de CARGA del ETL.

Contiene la clase ``Carga``, que:
    1. inserta los DataFrames transformados en una base de datos SQLite nueva,
    2. exporta los datos a archivos XLSX,
    3. verifica que los registros se hayan cargado correctamente,
    4. registra en el log los eventos principales.

Nota sobre Excel: una hoja admite como máximo 1.048.576 filas. Las tablas que
superan ese límite (p. ej. ``calendar``, con millones de filas) se dividen en
varias hojas o se excluyen del XLSX (siguen completas en SQLite y su versión
resumida ``calendar_mensual`` sí se exporta).
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from config import SQLITE_PATH, XLSX_DIR
from logger import obtener_logger

MAX_FILAS_EXCEL = 1_048_575  # 1.048.576 menos la fila de encabezados

# Índices para acelerar las consultas analíticas en SQLite
INDICES = {
    "listings": ["id", "host_id", "localidad", "room_type"],
    "hosts": ["host_id"],
    "listing_amenities": ["listing_id", "amenity"],
    "calendar": ["listing_id", "date", "anio_mes"],
    "calendar_mensual": ["listing_id", "anio_mes"],
    "reviews": ["listing_id", "date"],
}


class Carga:
    """Carga los datos transformados en SQLite y XLSX.

    Args:
        ruta_sqlite: archivo de la base de datos SQLite de destino.
        directorio_xlsx: carpeta donde se escriben los archivos Excel.
    """

    def __init__(self, ruta_sqlite: Path = SQLITE_PATH, directorio_xlsx: Path = XLSX_DIR) -> None:
        self.ruta_sqlite = Path(ruta_sqlite)
        self.directorio_xlsx = Path(directorio_xlsx)
        self.log = obtener_logger("carga")
        self.verificacion: list[dict] = []

    # ------------------------------------------------------------------ #
    # SQLite
    # ------------------------------------------------------------------ #
    @staticmethod
    def _preparar(df: pd.DataFrame) -> pd.DataFrame:
        """Convierte tipos que SQLite/Excel no manejan bien (category, listas)."""
        df = df.copy()
        for col in df.columns:
            if isinstance(df[col].dtype, pd.CategoricalDtype):
                df[col] = df[col].astype("string")
            elif df[col].dtype == object and df[col].map(lambda v: isinstance(v, (list, dict))).any():
                df[col] = df[col].astype(str)
        return df

    def cargar_sqlite(self, tablas: dict[str, pd.DataFrame], reemplazar_bd: bool = True) -> None:
        """Inserta cada DataFrame como una tabla en SQLite y crea índices.

        Args:
            tablas: diccionario {nombre_tabla: DataFrame}.
            reemplazar_bd: si es True, se crea una base de datos nueva en cada ejecución.
        """
        self.ruta_sqlite.parent.mkdir(parents=True, exist_ok=True)
        if reemplazar_bd and self.ruta_sqlite.exists():
            self.ruta_sqlite.unlink()
            self.log.info("Base de datos SQLite anterior eliminada: %s", self.ruta_sqlite.name)

        with sqlite3.connect(self.ruta_sqlite) as conexion:
            self.log.info("Conectado a SQLite: %s", self.ruta_sqlite)
            for nombre, df in tablas.items():
                try:
                    self._preparar(df).to_sql(nombre, conexion, if_exists="replace", index=False,
                                              chunksize=50_000)
                    self.log.info("Tabla %-18s cargada en SQLite: %s filas", nombre, f"{len(df):,}")
                    for col in INDICES.get(nombre, []):
                        if col in df.columns:
                            conexion.execute(
                                f'CREATE INDEX IF NOT EXISTS idx_{nombre}_{col} ON "{nombre}" ("{col}")')
                except Exception as error:  # noqa: BLE001
                    self.log.error("Error cargando la tabla %s en SQLite: %s", nombre, error)
                    raise
            conexion.commit()

    def verificar_sqlite(self, tablas: dict[str, pd.DataFrame]) -> bool:
        """Compara ``COUNT(*)`` de cada tabla en SQLite con las filas del DataFrame."""
        todo_ok = True
        with sqlite3.connect(self.ruta_sqlite) as conexion:
            for nombre, df in tablas.items():
                cargados = conexion.execute(f'SELECT COUNT(*) FROM "{nombre}"').fetchone()[0]
                ok = cargados == len(df)
                todo_ok &= ok
                self.verificacion.append({"destino": "SQLite", "tabla": nombre,
                                          "esperados": len(df), "cargados": cargados, "ok": ok})
                if ok:
                    self.log.info("Verificación SQLite OK  -> %-18s %s registros", nombre, f"{cargados:,}")
                else:
                    self.log.error("Verificación SQLite FALLÓ -> %s: esperados %s, cargados %s",
                                   nombre, len(df), cargados)
        return todo_ok

    # ------------------------------------------------------------------ #
    # XLSX
    # ------------------------------------------------------------------ #
    def exportar_xlsx(self, tablas: dict[str, pd.DataFrame], excluir: list[str] | None = None,
                      max_hojas: int = 3) -> dict[str, Path]:
        """Exporta cada tabla a su propio archivo ``<tabla>.xlsx``.

        Si una tabla supera el límite de filas de Excel se reparte en varias
        hojas (hasta ``max_hojas``); si aun así no cabe, se omite con WARNING.

        Args:
            tablas: diccionario {nombre: DataFrame}.
            excluir: tablas que no se exportan a Excel (p. ej. ``["calendar"]``).
            max_hojas: máximo de hojas por archivo para tablas grandes.

        Returns:
            Diccionario {nombre_tabla: ruta_del_archivo}.
        """
        excluir = set(excluir or [])
        self.directorio_xlsx.mkdir(parents=True, exist_ok=True)
        archivos: dict[str, Path] = {}

        for nombre, df in tablas.items():
            if nombre in excluir:
                self.log.warning("Tabla %s excluida del XLSX por tamaño (%s filas); disponible en SQLite.",
                                 nombre, f"{len(df):,}")
                continue
            partes = max(1, -(-len(df) // MAX_FILAS_EXCEL))  # división hacia arriba
            if partes > max_hojas:
                self.log.warning("Tabla %s (%s filas) requiere %d hojas de Excel; se omite del XLSX.",
                                 nombre, f"{len(df):,}", partes)
                continue
            ruta = self.directorio_xlsx / f"{nombre}.xlsx"
            try:
                datos = self._preparar(df)
                with pd.ExcelWriter(ruta, engine="openpyxl") as escritor:
                    for i in range(partes):
                        hoja = nombre[:28] if partes == 1 else f"{nombre[:25]}_{i + 1}"
                        datos.iloc[i * MAX_FILAS_EXCEL:(i + 1) * MAX_FILAS_EXCEL].to_excel(
                            escritor, sheet_name=hoja, index=False)
                archivos[nombre] = ruta
                self.log.info("Tabla %-18s exportada a %s (%d hoja/s)", nombre, ruta.name, partes)
            except Exception as error:  # noqa: BLE001
                self.log.error("Error exportando %s a XLSX: %s", nombre, error)
        return archivos

    def verificar_xlsx(self, archivos: dict[str, Path], tablas: dict[str, pd.DataFrame]) -> bool:
        """Cuenta las filas escritas en cada XLSX y las compara con el DataFrame."""
        todo_ok = True
        for nombre, ruta in archivos.items():
            libro = load_workbook(ruta, read_only=True)
            filas = sum(max(hoja.max_row - 1, 0) for hoja in libro.worksheets)
            libro.close()
            ok = filas == len(tablas[nombre])
            todo_ok &= ok
            self.verificacion.append({"destino": "XLSX", "tabla": nombre,
                                      "esperados": len(tablas[nombre]), "cargados": filas, "ok": ok})
            if ok:
                self.log.info("Verificación XLSX OK    -> %-18s %s registros", nombre, f"{filas:,}")
            else:
                self.log.error("Verificación XLSX FALLÓ -> %s: esperados %s, escritos %s",
                               nombre, len(tablas[nombre]), filas)
        return todo_ok

    # ------------------------------------------------------------------ #
    # Orquestación
    # ------------------------------------------------------------------ #
    def ejecutar(self, tablas: dict[str, pd.DataFrame], excluir_xlsx: list[str] | None = None) -> pd.DataFrame:
        """Carga en SQLite, exporta a XLSX y verifica ambos destinos.

        Returns:
            DataFrame con el resultado de la verificación por tabla y destino.
        """
        self.log.info("Iniciando carga de %d tablas", len(tablas))
        self.cargar_sqlite(tablas)
        ok_sqlite = self.verificar_sqlite(tablas)
        archivos = self.exportar_xlsx(tablas, excluir=excluir_xlsx)
        ok_xlsx = self.verificar_xlsx(archivos, tablas)
        if ok_sqlite and ok_xlsx:
            self.log.info("Carga finalizada correctamente: todos los conteos coinciden.")
        else:
            self.log.error("La carga terminó con diferencias en los conteos. Revise el log.")
        return pd.DataFrame(self.verificacion)
